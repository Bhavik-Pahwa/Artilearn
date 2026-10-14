"""Synthetic fixtures. The real gold set (PRD Sec. 12.1) is blocked on Q1 (subject choice)."""

from pathlib import Path

import pytest

PAGE_TEXT = "Normalization reduces redundancy in relations."


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid PDF; an empty string yields a text-less (image-only-like) page."""
    n = len(pages)
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        ("<< /Type /Pages /Kids [" + " ".join(f"{4 + 2 * i} 0 R" for i in range(n)) + f"] /Count {n} >>").encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % num + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


@pytest.fixture
def pdf_with_blank_page(tmp_path: Path) -> Path:
    p = tmp_path / "unit2.pdf"
    p.write_bytes(make_pdf([PAGE_TEXT, "", "Second normal form removes partial dependencies."]))
    return p


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("Unit 2", level=1)
    d.add_heading("Normalization", level=2)
    d.add_paragraph(PAGE_TEXT)
    d.add_heading("Unit 3", level=1)
    d.add_paragraph("Transactions follow ACID properties.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Atomicity"
    t.rows[0].cells[1].text = "All or nothing"
    p = tmp_path / "notes.docx"
    d.save(p)
    return p


@pytest.fixture
def pptx_with_empty_slide(tmp_path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    s1.shapes.title.text = "Normalization"
    s1.placeholders[1].text = PAGE_TEXT
    s1.notes_slide.notes_text_frame.text = "Mention 3NF next week."
    prs.slides.add_slide(prs.slide_layouts[6])  # blank: stands in for an image-only slide
    s3 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s3.shapes.title.text = "Functional dependencies"
    box = s3.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "A determines B if each A value maps to one B value."
    p = tmp_path / "slides.pptx"
    prs.save(p)
    return p
