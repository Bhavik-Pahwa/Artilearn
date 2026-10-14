"""Format-specific parsers (MAT-01, MAT-02). No OCR in the MVP: low-text units are flagged, not guessed.

Each parser returns a ParseResult. Locations follow the document schema:
- PDF:  page (1-based); heading_path unknown from the text layer, so left empty.
- DOCX: paragraph / table index in body order; heading_path from Heading N styles.
        DOCX has no fixed pages, so no page numbers are claimed.
- PPTX: slide (1-based); heading_path = [slide title]; speaker notes marked notes=True.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path


class ParseFailure(Exception):
    """Unparseable input. `code` is one of the schema's failure codes."""

    def __init__(self, code: str, detail: str):
        self.code = code
        self.detail = detail
        super().__init__(f"{code}: {detail}")


@dataclass
class RawBlock:
    text: str
    location: dict


@dataclass
class ParseResult:
    blocks: list[RawBlock] = field(default_factory=list)
    # One entry per page/slide (or one per document for DOCX): (location, chars)
    units: list[tuple[dict, int]] = field(default_factory=list)
    unit_count: int | None = None


def _paragraphs(text: str) -> list[str]:
    parts = [p.strip() for p in re.split(r"\n\s*\n", text)]
    return [p for p in parts if p]


def parse_pdf(path: Path) -> ParseResult:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ParseFailure("ENCRYPTED", "PDF is encrypted")
        result = ParseResult(unit_count=len(reader.pages))
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            result.units.append(({"page": i}, len(text.strip())))
            for para in _paragraphs(text):
                result.blocks.append(RawBlock(para, {"page": i, "heading_path": []}))
    except ParseFailure:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as e:
        raise ParseFailure("PARSE_ERROR", f"{type(e).__name__}: {e}") from e
    return result


def _heading_level(style_name: str | None) -> int | None:
    m = re.fullmatch(r"Heading (\d)", style_name or "")
    return int(m.group(1)) if m else None


def parse_docx(path: Path) -> ParseResult:
    import docx
    from docx.table import Table

    try:
        document = docx.Document(str(path))
    except Exception as e:  # python-docx raises a mix of zipfile/KeyError/ValueError
        raise ParseFailure("PARSE_ERROR", f"{type(e).__name__}: {e}") from e

    result = ParseResult(unit_count=None)
    heading_path: list[str] = []
    total = 0
    for idx, item in enumerate(document.iter_inner_content()):
        if isinstance(item, Table):
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in item.rows]
            text = "\n".join(r for r in rows if r.strip(" |"))
            loc = {"table": idx, "heading_path": list(heading_path)}
        else:
            text = item.text.strip()
            level = _heading_level(item.style.name if item.style is not None else None)
            if level is not None and text:
                heading_path = heading_path[: level - 1] + [text]
                total += len(text)
                continue
            loc = {"paragraph": idx, "heading_path": list(heading_path)}
        if text:
            total += len(text)
            result.blocks.append(RawBlock(text, loc))
    result.units.append(({"paragraph": 0}, total))
    return result


def parse_pptx(path: Path) -> ParseResult:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise ParseFailure("PARSE_ERROR", f"{type(e).__name__}: {e}") from e

    result = ParseResult(unit_count=len(prs.slides))
    for n, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        heading_path = [title] if title else []
        chars = 0
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
            elif getattr(shape, "has_table", False) and shape.has_table:
                text = "\n".join(" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows)
            else:
                continue
            if text:
                chars += len(text)
                path_for_block = [] if shape is title_shape else heading_path
                result.blocks.append(RawBlock(text, {"slide": n, "heading_path": path_for_block}))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                result.blocks.append(RawBlock(notes, {"slide": n, "notes": True, "heading_path": heading_path}))
        result.units.append(({"slide": n}, chars))
    return result


PARSERS = {".pdf": ("pdf", parse_pdf), ".docx": ("docx", parse_docx), ".pptx": ("pptx", parse_pptx)}
